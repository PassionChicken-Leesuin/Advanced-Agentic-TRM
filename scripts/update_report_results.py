# -*- coding: utf-8 -*-
"""사용자가 손본 진행보고 docx 를 다시 생성하지 않고, 전체 실행 결과가 들어갈 자리만 제자리에서 고친다.

python scripts/update_report_results.py --run runs/full [--docx <path>] [--out <path>] [--date "2026. 11. 1"]

고치는 곳 (앵커 = 현재 문서의 문장 시작/표 머리글):
  - [요약] 4번째 문단 (Initial/Advanced 비교)과 5번째 문단(비용)
  - 6.3 제목, <표 7> 캡션, 표 7 셀, 6.3 아래 불릿 2개와 → 결론 줄
  - <그림 1>/<그림 2> 캡션과 이미지(docs/screenshots/viz_kb_full.png, viz_roadmap_diff_full.png 가 있을 때)
  - <표 8> '전체 실행' 행과 '합계' 행
  - 9절 '180건 전체 실행 후 … 갱신할 예정임' 항목
서식은 각 문단의 첫 run 을 재사용해 유지한다.
"""
import argparse
import copy
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from docx import Document
from docx.shared import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_DOCX = os.path.join(HERE, "..", "..", "연구계획", "이수인_KB구축시스템_진행보고_1031.docx")
SCR = os.path.join(HERE, "..", "docs", "screenshots")


def pct(x):
    return "—" if x is None else f"{round(x * 100)}%"


def set_text(p, text, keep_bold_prefix=None):
    """문단의 첫 run 서식을 유지하며 텍스트 교체. keep_bold_prefix 가 있으면 그 접두는 굵게 둔 첫 run 에, 나머지는 둘째 run 에."""
    runs = p.runs
    if not runs:
        p.add_run(text); return
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    first = runs[0]
    if keep_bold_prefix and text.startswith(keep_bold_prefix):
        first.text = keep_bold_prefix
        second = copy.deepcopy(first._r); p._p.append(second)
        from docx.text.run import Run
        r2 = Run(second, p); r2.bold = None; r2.text = text[len(keep_bold_prefix):]
        rpr = second.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr")
        if rpr is not None:
            for b in rpr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}b"):
                rpr.remove(b)
    else:
        first.text = text


def set_cell(cell, text):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    set_text(p, text)


def find_para(doc, startswith=None, contains=None, after_idx=0):
    for i, p in enumerate(doc.paragraphs):
        if i < after_idx:
            continue
        t = p.text.strip()
        if (startswith and t.startswith(startswith)) or (contains and contains in t):
            return i, p
    return None, None


def replace_picture(doc, para, new_path):
    """문단 안의 첫 그림 blob 교체."""
    blips = para._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    if not blips:
        return False
    rid = blips[0].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
    part = doc.part.related_parts[rid]
    with open(new_path, "rb") as f:
        part._blob = f.read()
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    ap.add_argument("--docx", default=DEFAULT_DOCX)
    ap.add_argument("--out", default=None)
    ap.add_argument("--date", default=None)
    args = ap.parse_args()
    res = json.load(open(os.path.join(args.run, "summary.json"), encoding="utf-8"))
    kb, ai, aa = res["kb"], res["alignment"].get("initial", {}), res["alignment"].get("advanced", {})
    d, r = res["diff"], res["recovery"]; c = d["counts"]
    n_docs = kb["documents"]
    full_cost = res.get("full_cost_usd") or 0.0
    fin_cost = res.get("finalize_cost_usd") or 0.0
    conf = r.get("conflict_detection"); n_both = r.get("conflict_pairs_both_recovered", 0)
    conf_s = ("측정 불가(모순 쌍의 양쪽이 함께 회수된 경우 없음)" if conf is None else f"{pct(conf)} (양쪽 회수된 {n_both}쌍 기준)")
    better = (aa.get("recall_lenient") or 0) >= (ai.get("recall_lenient") or 0)
    mae_ok = aa.get("year_mae") is not None and ai.get("year_mae") is not None and aa["year_mae"] <= ai["year_mae"]

    doc = Document(args.docx)

    # ---- 날짜
    if args.date:
        for p in doc.paragraphs[:3]:
            if p.text.strip().startswith("DDSI Lab"):
                set_text(p, f"DDSI Lab 이수인 ({args.date})"); break

    # ---- [요약] 4·5번째 문단
    sp = doc.tables[0].rows[0].cells[0].paragraphs
    for p in sp:
        t = p.text
        if "Initial Roadmap" in t and "Advanced Roadmap" in t:
            set_text(p, (f"같은 LLM과 프롬프트를 사용해 KB 없이 만든 Initial Roadmap과 KB를 연산 ①~⑦로 활용해 만든 Advanced Roadmap을 비교함. "
                         f"{n_docs}건 전체 실행 결과 KB는 노드 {kb['nodes']}개, 주장 {kb['claims']}개, 수락 링크 {kb['edges_by_status'].get('accepted', 0)}개로 구성됨. "
                         f"가상 사실 회수율은 {pct(r.get('recovery_rate'))}(관대한 기준 {pct(r.get('recovery_rate_lenient'))})였음. "
                         f"Advanced는 {res['advanced']['nodes']}개 항목 중 {res['advanced']['with_evidence']}개에 근거 문서를 인용했고 Initial은 {res['initial']['with_evidence']}개였음. "
                         f"공식 로드맵 마일스톤 재현율(관대)은 Initial {pct(ai.get('recall_lenient'))}, Advanced {pct(aa.get('recall_lenient'))}이며, "
                         f"기준일 이후 공개된 2025년판 마일스톤의 사전 예측 재현율은 {pct(ai.get('recall_2025_lenient'))}에서 {pct(aa.get('recall_2025_lenient'))}로, "
                         f"연도 오차는 {ai.get('year_mae')}년에서 {aa.get('year_mae')}년으로 바뀜(6절)."))
        elif t.startswith("현재까지 비용"):
            total = 1.39 + 2.80 + 0.35 + full_cost + fin_cost
            set_text(p, f"전체 실행까지 포함한 비용은 약 ${total:.1f}임(전체 실행 {n_docs}건 ${full_cost:.2f}). 코드는 GitHub(PassionChicken-Leesuin/Advanced-Agentic-TRM)에 공개함.")

    # ---- 6.3 제목, 표 7 캡션
    _, p = find_para(doc, startswith="6.3 ")
    if p: set_text(p, "6.3 전체 실행의 KB 사용 전후 비교")
    _, p = find_para(doc, startswith="<표 7>")
    if p: set_text(p, f"<표 7> 전체 실행 결과 (KB 문서 {n_docs}건)")

    # ---- 표 7 셀
    t7 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "지표")
    rows = {t7.rows[i].cells[0].text.strip(): t7.rows[i] for i in range(1, len(t7.rows))}
    def setrow(label_key, v1, v2, new_label=None):
        row = next((rw for k, rw in rows.items() if k.startswith(label_key)), None)
        if row is None:   # 행 추가 (마지막 행 복제)
            tpl = t7.rows[-1]._tr; new = copy.deepcopy(tpl); t7._tbl.append(new)
            row = t7.rows[-1]; set_cell(row.cells[0], new_label or label_key)
        elif new_label:
            set_cell(row.cells[0], new_label)
        set_cell(row.cells[1], v1); set_cell(row.cells[2], v2)
    setrow("항목 수", str(res["initial"]["nodes"]), str(res["advanced"]["nodes"]))
    setrow("근거 문서를 인용한 항목", str(res["initial"]["with_evidence"]), str(res["advanced"]["with_evidence"]))
    setrow("가상 사실 회수율", "—", f"엄격 {pct(r.get('recovery_rate'))} / 관대 {pct(r.get('recovery_rate_lenient'))}")
    setrow("diff", "—", f"신규 {c.get('new', 0)}(근거 있음 {d['new_with_evidence']}) · 세분화 {c.get('refined', 0)} · 동일 {c.get('same', 0)} · 일반화 {c.get('generalized', 0)} · 탈락 {c.get('dropped', 0)}")
    setrow("공식 로드맵 recall", f"{pct(ai.get('recall_strict'))} / {pct(ai.get('recall_lenient'))}", f"{pct(aa.get('recall_strict'))} / {pct(aa.get('recall_lenient'))}")
    setrow("2025년판 recall", pct(ai.get("recall_2025_lenient")), pct(aa.get("recall_2025_lenient")))
    setrow("연도 오차 MAE", f"{ai.get('year_mae')}년", f"{aa.get('year_mae')}년")
    setrow("모순 탐지율", "—", conf_s, new_label="모순 탐지율 (모순 쌍 3개)")
    setrow("2020년판 recall", pct(ai.get("recall_2020_lenient")), pct(aa.get("recall_2020_lenient")), new_label="2020년판 recall (관대)")

    # ---- 6.3 불릿 2개 + → 결론
    i, p = find_para(doc, startswith="문서 20건으로 진행한 파일럿")
    if p:
        set_text(p, (f"문서 {n_docs}건 전체 실행에서 Advanced는 {res['advanced']['nodes']}개 항목 중 {res['advanced']['with_evidence']}개에 근거 문서를 인용했고(Initial {res['initial']['with_evidence']}개), "
                     f"연도 오차는 {ai.get('year_mae')}년에서 {aa.get('year_mae')}년으로 " + ("줄었음. " if mae_ok else "줄지 않았음. ")
                     + f"공식 마일스톤 재현율(관대)은 Initial {pct(ai.get('recall_lenient'))}, Advanced {pct(aa.get('recall_lenient'))}로 "
                     + ("Advanced가 같거나 높았음. " if better else "Initial이 아직 높았음. ")
                     + f"항목 변화는 신규 {c.get('new', 0)}건(근거 있음 {d['new_with_evidence']}건), 세분화 {c.get('refined', 0)}건, 탈락 {c.get('dropped', 0)}건임."))
    i, p = find_para(doc, startswith="파일럿에서는 모순 탐지율")
    if p:
        set_text(p, "모순 탐지율은 사전에 정한 모순 쌍 3개 가운데 양쪽 사실이 모두 KB에 회수된 쌍에서 Critic이 conflict를 표시한 비율임. 전체 실행 결과: " + conf_s + ".")
    i, p = find_para(doc, startswith="→ 공식 로드맵과의 비교는")
    if p:
        set_text(p, "→ 전체 실행에서 확인된 KB의 효과는 근거가 붙은 항목 수와 시점 정확도에서 가장 뚜렷함. 공식 마일스톤 재현율은 "
                 + ("Advanced가 Initial과 같거나 높아, 20건 파일럿에서 보였던 내부자료 편중이 문서 수가 늘면서 완화됨." if better
                    else "아직 Initial보다 낮아, 연산 ⑦의 우선순위를 조정해 공개 자료의 주요 마일스톤이 내부자료의 세부 항목보다 먼저 반영되도록 보완할 계획임."))

    # ---- 그림 교체·캡션
    kb_png, rm_png = os.path.join(SCR, "viz_kb_full.png"), os.path.join(SCR, "viz_roadmap_diff_full.png")
    i1, cap1 = find_para(doc, startswith="<그림 1>")
    if cap1 and os.path.exists(kb_png):
        for j in range(i1 - 1, max(0, i1 - 3), -1):
            if replace_picture(doc, doc.paragraphs[j], kb_png):
                set_text(cap1, f"<그림 1> 2막 지식베이스 구축 화면 (전체 실행, 문서 {n_docs}건 처리 후)"); break
    i2, cap2 = find_para(doc, startswith="<그림 2>")
    if cap2 and os.path.exists(rm_png):
        for j in range(i2 - 1, max(0, i2 - 3), -1):
            if replace_picture(doc, doc.paragraphs[j], rm_png):
                set_text(cap2, "<그림 2> Initial(왼쪽)·Advanced(오른쪽) 로드맵과 항목 변화 표시 (전체 실행)"); break

    # ---- 표 8
    t8 = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "항목")
    for rw in t8.rows[1:]:
        lab = rw.cells[0].text.strip()
        if lab.startswith("전체 실행"):
            set_cell(rw.cells[0], f"전체 실행 {n_docs}건 (완료, 문서당 약 ${full_cost / max(1, n_docs):.3f})"); set_cell(rw.cells[1], f"${full_cost:.2f}")
        elif lab.startswith("합계"):
            set_cell(rw.cells[0], "합계"); set_cell(rw.cells[1], f"약 ${1.39 + 2.80 + 0.35 + full_cost + fin_cost:.1f}")
        elif lab.startswith("로드맵 생성"):
            set_cell(rw.cells[1], f"약 ${0.35 + fin_cost:.2f}")

    # ---- 9절 첫 항목
    i, p = find_para(doc, startswith="180건 전체 실행 후")
    if p:
        set_text(p, f"{n_docs}건 전체 실행 결과를 <표 7>과 요약에 반영함. 남은 작업은 결과 해석과 본 실험 설계(지식 구조 조건별 비교)로의 연결임.")

    out = args.out or args.docx
    doc.save(out)
    print("updated", out)


if __name__ == "__main__":
    main()
